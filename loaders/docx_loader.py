from pathlib import Path
import docx
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document


def _iter_block_items(document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _table_to_markdown(table: Table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header, *body = rows
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def load_docx(file_path: str) -> list[Document]:
    filename = Path(file_path).name
    document = docx.Document(file_path)

    documents: list[Document] = []
    text_buffer: list[str] = []
    section_index = 1
    table_index = 1

    def flush_text_buffer():
        nonlocal section_index
        content = "\n".join(text_buffer).strip()
        if content:
            documents.append(Document(
                page_content=content,
                metadata={"source_filename": filename, "section_number": section_index, "type": "text"},
            ))
            section_index += 1
        text_buffer.clear()

    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            if block.text.strip():
                text_buffer.append(block.text.strip())
        elif isinstance(block, Table):
            flush_text_buffer()  # tables are their own chunk, never merged with prose
            markdown = _table_to_markdown(block)
            if markdown.strip():
                documents.append(Document(
                    page_content=markdown,
                    metadata={"source_filename": filename, "table_index": table_index, "type": "table"},
                ))
                table_index += 1

    flush_text_buffer()
    return documents